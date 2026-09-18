"""The web preview must show exactly the document the .docx export writes."""

from __future__ import annotations

import io
import uuid

import docx
import pytest
from httpx import ASGITransport, AsyncClient

from app.export.docx_renderer import render_docx
from app.export.preview import render_preview
from app.export.profiles import ExportOptions, RenderProfile
from app.main import app
from app.resume.parser import parse


def _block_text(block) -> str:
    if block.kind == "role":
        return f"{block.text}\t{block.right}" if block.right else block.text
    if block.kind == "bullet":
        return f"•\t{block.text}"
    if block.kind == "skill_line":
        return f"{block.label}: {block.text}"
    return block.text


@pytest.mark.parametrize("profile", list(RenderProfile))
def test_blocks_match_docx_paragraphs(aran_markdown: str, profile) -> None:
    resume = parse(aran_markdown)
    options = ExportOptions(profile=profile)
    document = docx.Document(io.BytesIO(render_docx(resume, options)))
    preview = render_preview(resume, options)

    # doc.paragraphs excludes table cells, so the skills table is compared apart.
    assert [(p.style.name, p.text) for p in document.paragraphs] == [
        (b.style, _block_text(b)) for b in preview.blocks if b.kind != "skills_table"
    ]
    table_rows = [
        (row.label, row.value)
        for b in preview.blocks
        if b.kind == "skills_table"
        for row in b.rows
    ]
    assert table_rows == [
        (r.cells[0].text, r.cells[1].text) for t in document.tables for r in t.rows
    ]


def test_styles_come_from_the_carrier(aran_markdown: str) -> None:
    resume = parse(aran_markdown)
    designed = render_preview(resume, ExportOptions(profile=RenderProfile.DESIGNED))
    ats = render_preview(resume, ExportOptions(profile=RenderProfile.ATS_PLAIN))

    section = designed.styles["ResumeSection"]
    assert section.bold and section.all_caps and section.border_bottom is not None
    assert section.border_bottom.width_pt == 0.75
    assert designed.styles["ResumeSummary"].justify
    assert designed.styles["ResumeName"].color == "1F3864"
    assert designed.layout.page_width_pt == pytest.approx(612, abs=0.1)
    assert designed.layout.margin_left_pt == pytest.approx(0.65 * 72, abs=0.1)

    assert ats.styles["ResumeSection"].border_bottom is None
    assert not ats.styles["ResumeSummary"].justify
    assert ats.styles["ResumeBody"].size_pt == 11


@pytest.fixture
async def client():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


async def test_preview_endpoint_accepts_unsaved_markdown(client, aran_markdown: str) -> None:
    res = await client.post(
        "/api/v1/auth/register",
        json={"email": f"preview-{uuid.uuid4().hex[:8]}@example.com", "password": "a-strong-password"},
    )
    headers = {"Authorization": f"Bearer {res.json()['access_token']}"}

    res = await client.post(
        "/api/v1/export/preview", headers=headers, json={"content_markdown": aran_markdown}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["blocks"][0] == {
        "kind": "paragraph",
        "style": "ResumeName",
        "text": parse(aran_markdown).header.name,
        "right": None,
        "label": None,
        "rows": None,
    }

    res = await client.post(
        "/api/v1/export/preview", headers=headers, json={"content_markdown": "not a resume"}
    )
    assert res.status_code == 422
