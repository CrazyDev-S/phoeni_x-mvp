"""Reading a resume out of an uploaded document.

Resumes arrive as PDFs, Word files, OpenDocument files, RTF, HTML and plain
text. Everything is reduced to the same thing: the document's text in reading
order, with list items marked and hyperlink targets kept, ready for the model
to restructure. Nothing here interprets the content.

Uploads are untrusted. Sizes are capped before anything is opened, archives are
checked for decompression bombs, and XML is parsed without entities or network
access.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass, field

from lxml import etree

MAX_BYTES = 8 * 1024 * 1024
MAX_PDF_PAGES = 30
# Far above any resume; past this the file is something else.
MAX_CHARS = 120_000
MIN_CHARS = 80
# Uncompressed size of an OOXML/ODF package, to refuse a decompression bomb.
MAX_UNZIPPED = 60 * 1024 * 1024

log = logging.getLogger(__name__)

SUPPORTED = "PDF, Word (.docx), OpenDocument (.odt), RTF, HTML, Markdown or plain text"


class DocumentError(ValueError):
    """The file cannot be read as a resume. The message is for the user."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ExtractedDocument:
    text: str
    kind: str
    pages: int | None = None
    # Hyperlink targets found in the file; the visible text often omits them.
    links: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def for_model(self) -> str:
        """The text with its links appended, where the model can see them."""
        if not self.links:
            return self.text
        return self.text + "\n\nLINKS FOUND IN THE FILE:\n" + "\n".join(f"- {u}" for u in self.links)


def extract(filename: str, data: bytes) -> ExtractedDocument:
    """Read any supported file. Raises DocumentError with an actionable message."""
    if not data:
        raise DocumentError("EMPTY_FILE", "The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(
            "FILE_TOO_LARGE",
            f"The file is {len(data) / 1_048_576:.1f} MB; the limit is {MAX_BYTES // 1_048_576} MB.",
        )

    try:
        doc = _read(filename, data)
    except DocumentError:
        raise
    except Exception as exc:
        # Parsers meet hostile and damaged bytes: pypdf, zipfile, zlib and lxml
        # each raise their own types for them. Whatever it is, the caller needs
        # a message, not a 500 - which a browser reports as "Failed to fetch".
        log.warning("could not read %r: %s: %s", filename, type(exc).__name__, exc)
        raise DocumentError(
            "CORRUPT_FILE",
            "This file could not be read - it may be damaged. Export it again, "
            "or upload it as a different format.",
        ) from exc

    if len(doc.text) < MIN_CHARS:
        if doc.kind == "pdf":
            raise DocumentError(
                "NO_TEXT_LAYER",
                "This PDF has no readable text - it looks like a scan or an image. "
                "Export it again from the original program, or upload a Word file.",
            )
        raise DocumentError("NO_TEXT", "There is almost no text in this file to read.")
    if len(doc.text) > MAX_CHARS:
        doc.text = doc.text[:MAX_CHARS]
        doc.warnings.append("The document is very long; only the first part was read.")
    return doc


def _read(filename: str, data: bytes) -> ExtractedDocument:
    kind = sniff(filename, data)
    reader = {
        "pdf": _read_pdf,
        "docx": _read_docx,
        "odt": _read_odt,
        "rtf": _read_rtf,
        "html": _read_html,
        "text": _read_text,
    }[kind]
    doc = reader(data)
    doc.kind = kind
    doc.text = _tidy(doc.text)
    return doc


def sniff(filename: str, data: bytes) -> str:
    """The format from the bytes, falling back on the extension for text.

    The extension alone is not trusted: a renamed file would otherwise be
    handed to the wrong reader.
    """
    head = data[:8]
    name = filename.lower()
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        raise DocumentError(
            "LEGACY_DOC",
            "Old binary Word (.doc) files are not supported. Save it as .docx or PDF and upload that.",
        )
    if head.startswith(b"PK\x03\x04"):
        return _sniff_package(data)
    if data[:512].count(b"\x00") > 8:
        raise DocumentError(
            "UNSUPPORTED_TYPE",
            f"This file type is not supported. Upload {SUPPORTED}.",
        )
    if name.endswith((".html", ".htm")) or re.search(rb"<\s*(html|body|div|p)\b", data[:2048], re.I):
        return "html"
    return "text"


def _sniff_package(data: bytes) -> str:
    try:
        archive = _open_zip(data)
    except zipfile.BadZipFile as exc:
        raise DocumentError("CORRUPT_FILE", "The file is damaged and could not be opened.") from exc
    names = set(archive.namelist())
    if "word/document.xml" in names:
        return "docx"
    if "content.xml" in names and "mimetype" in names:
        return "odt"
    raise DocumentError(
        "UNSUPPORTED_TYPE", f"This file type is not supported. Upload {SUPPORTED}."
    )


def _open_zip(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    infos = archive.infolist()
    if sum(i.file_size for i in infos) > MAX_UNZIPPED or any(
        i.file_size > MAX_UNZIPPED // 2 for i in infos
    ):
        raise DocumentError("FILE_TOO_LARGE", "The document expands to an unreasonable size.")
    return archive


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _read_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    # pypdf narrates every repair it makes to a damaged file; that is noise here.
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError(
                "PASSWORD_PROTECTED",
                "This PDF is password protected. Remove the password and upload it again.",
            )
        pages = reader.pages
        count = len(pages)
        parts: list[str] = []
        links: list[str] = []
        for page in list(pages)[:MAX_PDF_PAGES]:
            parts.append(page.extract_text(extraction_mode="layout") or "")
            links.extend(_pdf_links(page))
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, AttributeError) as exc:
        raise DocumentError("CORRUPT_FILE", "The PDF is damaged and could not be read.") from exc

    doc = ExtractedDocument(text="\n\n".join(parts), kind="pdf", pages=count, links=_unique(links))
    if count > MAX_PDF_PAGES:
        doc.warnings.append(f"Only the first {MAX_PDF_PAGES} of {count} pages were read.")
    return doc


def _pdf_links(page) -> list[str]:
    found: list[str] = []
    for annot in page.get("/Annots") or []:
        try:
            action = annot.get_object().get("/A")
            uri = action.get_object().get("/URI") if action else None
        except (AttributeError, KeyError, ValueError):
            continue
        if isinstance(uri, str) and uri.startswith(("http://", "https://", "mailto:")):
            found.append(uri)
    return found


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _read_docx(data: bytes) -> ExtractedDocument:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    _open_zip(data)
    try:
        document = Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise DocumentError("CORRUPT_FILE", "The Word file is damaged and could not be opened.") from exc

    rels = document.part.rels
    links: list[str] = []
    lines: list[str] = []

    # Contact details very often live in the page header.
    seen_headers: set[int] = set()
    for section in document.sections:
        header = section.header
        if id(header._element) in seen_headers or header.is_linked_to_previous:
            continue
        seen_headers.add(id(header._element))
        lines.extend(_word_paragraphs(header._element, rels, links))

    # Every paragraph in document order, which takes in tables and text boxes.
    lines.extend(_word_paragraphs(document.element.body, rels, links))
    return ExtractedDocument(text="\n".join(lines), kind="docx", links=_unique(links))


def _word_paragraphs(root, rels, links: list[str]) -> list[str]:
    out: list[str] = []
    for p in root.iter(f"{{{_W}}}p"):
        # A text box is written twice - once modern, once as a fallback.
        if any(a.tag == f"{{{_MC}}}Fallback" for a in p.iterancestors()):
            continue
        parts: list[str] = []
        for node in p.iter():
            tag = etree.QName(node).localname if isinstance(node.tag, str) else ""
            if tag == "t":
                parts.append(node.text or "")
            elif tag == "tab":
                parts.append("\t")
            elif tag in ("br", "cr"):
                parts.append("\n")
            elif tag == "hyperlink":
                rid = node.get(f"{{{_R}}}id")
                target = rels[rid].target_ref if rid and rid in rels else None
                if isinstance(target, str) and target.startswith(("http://", "https://", "mailto:")):
                    links.append(target)
        text = "".join(parts).strip()
        if not text:
            continue
        marked = p.find(f"{{{_W}}}pPr/{{{_W}}}numPr") is not None
        out.append(f"• {text}" if marked else text)
    return out


# --------------------------------------------------------------------------
# OpenDocument
# --------------------------------------------------------------------------

_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_XLINK = "http://www.w3.org/1999/xlink"


def _read_odt(data: bytes) -> ExtractedDocument:
    archive = _open_zip(data)
    try:
        parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
        root = etree.fromstring(archive.read("content.xml"), parser)
    except (KeyError, etree.XMLSyntaxError) as exc:
        raise DocumentError("CORRUPT_FILE", "The OpenDocument file could not be read.") from exc

    links = [
        h
        for a in root.iter(f"{{{_TEXT}}}a")
        if (h := a.get(f"{{{_XLINK}}}href", "")).startswith(("http://", "https://", "mailto:"))
    ]
    lines: list[str] = []
    for node in root.iter(f"{{{_TEXT}}}p", f"{{{_TEXT}}}h"):
        text = _odf_text(node).strip()
        if not text:
            continue
        in_list = any(a.tag == f"{{{_TEXT}}}list-item" for a in node.iterancestors())
        lines.append(f"• {text}" if in_list else text)
    return ExtractedDocument(text="\n".join(lines), kind="odt", links=_unique(links))


def _odf_text(node) -> str:
    parts = [node.text or ""]
    for child in node:
        tag = etree.QName(child).localname if isinstance(child.tag, str) else ""
        if tag == "s":
            parts.append(" " * int(child.get(f"{{{_TEXT}}}c", "1")))
        elif tag == "tab":
            parts.append("\t")
        elif tag == "line-break":
            parts.append("\n")
        else:
            parts.append(_odf_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


# --------------------------------------------------------------------------
# RTF, HTML, text
# --------------------------------------------------------------------------

_RTF_SKIP = {
    "fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer",
    "themedata", "colorschememapping", "latentstyles", "datastore", "listtable",
    "listoverridetable", "rsidtbl", "generator", "filetbl", "revtbl",
}
_RTF_TOKEN = re.compile(
    r"\\(?P<word>[a-zA-Z]+)(?P<num>-?\d+)? ?|\\'(?P<hex>[0-9a-fA-F]{2})|\\(?P<sym>[^a-zA-Z])|(?P<open>\{)|(?P<close>\})|(?P<text>[^\\{}]+)"
)


def _read_rtf(data: bytes) -> ExtractedDocument:
    """A small RTF reader: enough for text in order, not for layout."""
    source = data.decode("latin-1")
    out: list[str] = []
    stack: list[bool] = []
    skipping = False
    pending_ignorable = False
    # \uN is followed by one plain-ASCII stand-in for readers without Unicode.
    fallback = 0

    for m in _RTF_TOKEN.finditer(source):
        if m.group("open"):
            stack.append(skipping)
        elif m.group("close"):
            if stack:
                skipping = stack.pop()
            pending_ignorable = False
        elif m.group("word"):
            word, num = m.group("word"), m.group("num")
            if (word in _RTF_SKIP or pending_ignorable) and not skipping:
                skipping = True
            elif skipping:
                pass
            elif word in ("par", "line", "row", "sect", "page"):
                out.append("\n")
            elif word == "tab":
                out.append("\t")
            elif word == "bullet":
                out.append("•")
            elif word == "endash":
                out.append("–")
            elif word == "emdash":
                out.append("—")
            elif word == "u" and num:
                code = int(num)
                out.append(chr(code + 65536 if code < 0 else code))
                fallback = 1
            pending_ignorable = False
        elif m.group("sym") is not None:
            sym = m.group("sym")
            if sym == "*":
                pending_ignorable = True
            elif sym in "\\{}" and not skipping:
                out.append(sym)
            elif sym == "~" and not skipping:
                out.append(" ")
        elif m.group("hex"):
            if fallback:
                fallback = 0
            elif not skipping:
                out.append(bytes([int(m.group("hex"), 16)]).decode("cp1252", "replace"))
        elif m.group("text") and not skipping:
            text = m.group("text").replace("\r", "").replace("\n", "")
            out.append(text[fallback:])
            fallback = 0
    return ExtractedDocument(text="".join(out), kind="rtf")


_BLOCKS = ["p", "div", "li", "tr", "section", "article", "header", "footer", "ul", "ol", "table", "h1", "h2", "h3", "h4", "h5", "h6"]


def _read_html(data: bytes) -> ExtractedDocument:
    from bs4 import BeautifulSoup, Comment, Doctype

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "head"]):
        tag.decompose()
    links = [
        a["href"]
        for a in soup.find_all("a", href=True)
        if str(a["href"]).startswith(("http://", "https://", "mailto:"))
    ]
    # Text wraps wherever the source did; only block elements start a new line.
    for node in soup.find_all(string=True):
        if isinstance(node, Comment | Doctype):
            node.extract()
        elif not node.find_parent("pre"):
            node.replace_with(re.sub(r"\s+", " ", str(node)))
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for li in soup.find_all("li"):
        li.insert(0, "• ")
    for block in soup.find_all(_BLOCKS):
        block.insert_before("\n")
        block.insert_after("\n")
    return ExtractedDocument(
        text=re.sub(r"\n[ \t]+", "\n", soup.get_text("")),
        kind="html",
        links=_unique([str(h) for h in links]),
    )


def _read_text(data: bytes) -> ExtractedDocument:
    return ExtractedDocument(text=_decode(data), kind="text")


def _decode(data: bytes) -> str:
    # UTF-16 only on its byte-order mark: without one, nearly any even-length
    # byte string "decodes", as garbage.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "replace")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


# --------------------------------------------------------------------------

def _unique(items: list[str]) -> list[str]:
    seen: dict[str, None] = {}
    for item in items:
        seen.setdefault(item.strip(), None)
    return [i for i in seen if i][:30]


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    # Layout extraction pads columns with long runs of spaces.
    text = re.sub(r"(?<=\S) {3,}(?=\S)", "   ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
